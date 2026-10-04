import json
from pathlib import Path

from openpyxl import load_workbook
from typer.testing import CliRunner

from reconcile_skill.cli import app


def test_cli_end_to_end_and_mapping_rerun(tmp_path: Path):
    source = tmp_path / "source.csv"
    target = tmp_path / "target.csv"
    source.write_text("id,amount,name\nA001,100,Alice\nA002,200,Bob\n", encoding="utf-8")
    target.write_text("id,amount,name\nA001,100,Alice\nA002,201,Bob\n", encoding="utf-8")
    first_output = tmp_path / "first"
    second_output = tmp_path / "second"
    runner = CliRunner()

    first = runner.invoke(
        app,
        ["--source", str(source), "--target", str(target), "--output", str(first_output)],
        input="1\n1\n",
    )
    assert first.exit_code == 0, first.stdout
    for filename in ("reconciliation_report.xlsx", "mapping.yaml", "run_log.json"):
        assert (first_output / filename).exists()

    second = runner.invoke(
        app,
        [
            "--source",
            str(source),
            "--target",
            str(target),
            "--output",
            str(second_output),
            "--mapping",
            str(first_output / "mapping.yaml"),
            "--non-interactive",
        ],
    )
    assert second.exit_code == 0, second.stdout
    first_log = json.loads((first_output / "run_log.json").read_text(encoding="utf-8"))
    second_log = json.loads((second_output / "run_log.json").read_text(encoding="utf-8"))
    assert first_log["summary"] == second_log["summary"]
    assert first_log["summary"]["value_mismatches"] == 1

    def identity_rows(report: Path, sheet: str, fields: tuple[str, ...]):
        workbook = load_workbook(report, read_only=True, data_only=True)
        rows = workbook[sheet].iter_rows(values_only=True)
        try:
            headers = next(rows)
        except StopIteration:
            return set()
        positions = {header: index for index, header in enumerate(headers)}
        return {
            tuple(row[positions[field]] for field in fields)
            for row in rows
        }

    assert identity_rows(
        first_output / "reconciliation_report.xlsx",
        "Missing_In_Source",
        ("reconciliation_key",),
    ) == identity_rows(
        second_output / "reconciliation_report.xlsx",
        "Missing_In_Source",
        ("reconciliation_key",),
    )
    assert identity_rows(
        first_output / "reconciliation_report.xlsx",
        "Missing_In_Target",
        ("reconciliation_key",),
    ) == identity_rows(
        second_output / "reconciliation_report.xlsx",
        "Missing_In_Target",
        ("reconciliation_key",),
    )
    assert identity_rows(
        first_output / "reconciliation_report.xlsx",
        "Value_Mismatch",
        ("reconciliation_key", "source_column", "target_column", "difference_type"),
    ) == identity_rows(
        second_output / "reconciliation_report.xlsx",
        "Value_Mismatch",
        ("reconciliation_key", "source_column", "target_column", "difference_type"),
    )
