from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from reconcile_skill.models import CompareRule
from reconcile_skill.normalizers import parse_datetime, parse_number, values_equal
from reconcile_skill.reconciler import reconcile
from reconcile_skill.report import write_report


@pytest.mark.parametrize(
    ("value", "expected"),
    [
        ("1,234.50", "1234.50"),
        ("-1,234", "-1234"),
        ("1e3", "1000"),
        (".5", "0.5"),
    ],
)
def test_strict_number_parser_accepts_only_valid_numeric_syntax(value, expected):
    assert parse_number(value) == parse_number(expected)


@pytest.mark.parametrize("value", ["100USD", "100kg", "100%", "12,34", "abc100", "NaN", "Infinity"])
def test_strict_number_parser_rejects_units_text_invalid_grouping_and_non_finite(value):
    assert parse_number(value) is None


def test_boolean_is_not_a_numeric_value():
    equal, difference_type = values_equal(True, 1, "numeric", None, ["numeric_parse"])
    assert equal is False
    assert difference_type.startswith("UNVERIFIED")


@pytest.mark.parametrize("tolerance", [-0.01, float("nan"), float("inf")])
def test_invalid_tolerance_is_rejected(tolerance):
    with pytest.raises((ValueError, TypeError)):
        CompareRule(source_column="amount", target_column="amount", rule_type="numeric", tolerance=tolerance)


def test_same_invalid_numeric_text_is_not_a_match():
    equal, difference_type = values_equal("abc100", "abc100", "numeric", None, ["numeric_parse"])
    assert equal is False
    assert difference_type == "UNVERIFIED_NUMERIC_PARSE"


def test_invalid_numeric_against_null_is_unverified_not_null_mismatch():
    equal, difference_type = values_equal(None, "abc100", "numeric", None, ["numeric_parse"])
    assert equal is False
    assert difference_type == "UNVERIFIED_NUMERIC_PARSE"


def test_numeric_units_and_percent_are_not_silently_stripped():
    for left, right in [("100USD", "100EUR"), ("100kg", "100g"), ("100%", "100")]:
        equal, difference_type = values_equal(left, right, "numeric", None, ["numeric_parse"])
        assert equal is False
        assert difference_type == "UNVERIFIED_NUMERIC_PARSE"


def test_nan_and_infinity_are_not_normal_numbers():
    for left, right in [("NaN", 100), ("Infinity", "Infinity")]:
        equal, difference_type = values_equal(left, right, "numeric", None, ["numeric_parse"])
        assert equal is False
        assert difference_type == "UNVERIFIED_NON_FINITE"


def test_datetime_parse_failure_does_not_fallback_to_string_match():
    assert parse_datetime("2026-02-30") is None
    equal, difference_type = values_equal("not-a-date", "not-a-date", "datetime", None, ["datetime_parse"])
    assert equal is False
    assert difference_type == "UNVERIFIED_DATETIME_PARSE"


def test_mixed_naive_and_timezone_datetime_is_unverified():
    equal, difference_type = values_equal(
        "2026-01-01T00:00:00",
        "2026-01-01T00:00:00+00:00",
        "datetime",
        None,
        ["datetime_parse"],
    )
    assert equal is False
    assert difference_type == "UNVERIFIED_TIMEZONE_MISSING"


def test_unverified_value_is_separate_from_mismatch_and_match():
    result = reconcile(
        pd.DataFrame({"id": ["A1"], "amount": ["abc100"]}),
        pd.DataFrame({"id": ["A1"], "amount": ["abc100"]}),
        source_key="id",
        target_key="id",
        rules=[CompareRule(source_column="amount", target_column="amount", rule_type="numeric")],
    )
    assert result.summary.value_mismatch_count == 0
    assert result.summary.unverified_count == 1
    assert result.unverified[0]["difference_type"] == "UNVERIFIED_NUMERIC_PARSE"
    assert result.status == "PARTIAL_NEEDS_REVIEW"
    assert any("无法可靠比较" in warning for warning in result.warnings)


def test_empty_rules_are_not_reported_as_complete_reconciliation():
    result = reconcile(
        pd.DataFrame({"id": ["A1"], "name": ["Alice"]}),
        pd.DataFrame({"id": ["A1"], "name": ["Alice"]}),
        source_key="id",
        target_key="id",
        rules=[],
    )
    assert result.summary.value_mismatch_count == 0
    assert result.summary.compared_field_count == 0
    assert result.summary.coverage_status == "PARTIAL_NEEDS_REVIEW"
    assert result.status == "PARTIAL_NEEDS_REVIEW"


def test_field_coverage_exposes_pending_and_skipped_mappings():
    result = reconcile(
        pd.DataFrame({"id": ["A1"], "name": ["Alice"], "status": ["paid"]}),
        pd.DataFrame({"id": ["A1"], "full_name": ["Alice"], "state": ["SUCCESS"]}),
        source_key="id",
        target_key="id",
        rules=[CompareRule(source_column="id", target_column="id")],
        field_coverage=[
            {"source_column": "id", "target_column": "id", "status": "COMPARED"},
            {"source_column": "name", "target_column": "full_name", "status": "PENDING_CONFIRMATION"},
            {"source_column": "status", "target_column": "state", "status": "SKIPPED"},
        ],
    )
    assert result.summary.pending_field_count == 1
    assert result.summary.skipped_field_count == 1
    assert result.summary.compared_field_count == 1
    assert result.status == "PARTIAL_NEEDS_REVIEW"


def test_missing_diagnostics_cover_both_directions():
    result = reconcile(
        pd.DataFrame({"id": ["S1", "S2"], "status": ["OPEN", "OPEN"]}),
        pd.DataFrame({"id": ["S1", "T1"], "status": ["OPEN", "CLOSED"]}),
        source_key="id",
        target_key="id",
        rules=[],
    )
    directions = {item["direction"] for item in result.diagnostics}
    assert directions == {"Missing_In_Source", "Missing_In_Target"}
    assert all("root cause" not in item["hypothesis"].lower() for item in result.diagnostics)


def test_report_exposes_unverified_and_field_coverage(tmp_path: Path):
    result = reconcile(
        pd.DataFrame({"id": ["A1"], "amount": ["100USD"]}),
        pd.DataFrame({"id": ["A1"], "amount": ["100EUR"]}),
        source_key="id",
        target_key="id",
        rules=[CompareRule(source_column="amount", target_column="amount", rule_type="numeric")],
    )
    report, _ = write_report(result, tmp_path, source_key="id", target_key="id")
    workbook = load_workbook(report, read_only=True)
    assert "Unverified" in workbook.sheetnames
    summary = {row[0].value: row[1].value for row in workbook["Summary"].iter_rows() if row[0].value}
    assert summary["Unverified Count"] == 1
    assert summary["Reconciliation Status"] == "PARTIAL_NEEDS_REVIEW"
