from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from reconcile_skill.errors import ReconciliationError
from reconcile_skill.loaders import load_table
from reconcile_skill.mapping import rules_from_suggestions
from reconcile_skill.models import MappingSuggestion
from reconcile_skill.models import CompareRule
from reconcile_skill.report import sanitize_excel_cell, write_report
from reconcile_skill.semantic_matcher import suggest_column_mappings
from reconcile_skill.loaders import LoadedTable
from reconcile_skill.profiler import profile_dataframe


def test_chinese_mapping_golden_set_and_wrong_key_trap():
    source = pd.DataFrame(
        {
            "订单号": ["A001"],
            "客户编号": ["C001"],
            "客户名称": ["Alice"],
            "手机号": ["13800000000"],
            "订单金额": [100],
            "实付金额": [98],
            "订单状态": ["paid"],
            "创建时间": ["2026-01-01"],
            "商品编码": ["SKU-1"],
        }
    )
    target = pd.DataFrame(
        {
            "order_id": ["A001"],
            "customer_id": ["C001"],
            "customer_name": ["Alice"],
            "mobile": ["13800000000"],
            "total_amount": [100],
            "paid_amount": [98],
            "order_status": ["SUCCESS"],
            "created_at": ["2026-01-01"],
            "sku": ["SKU-1"],
        }
    )
    left, right = LoadedTable(source, "source.csv", None), LoadedTable(target, "target.csv", None)
    suggestions = suggest_column_mappings(
        profile_dataframe(left), profile_dataframe(right), source, target
    )
    pairs = {(item.source_column, item.target_column) for item in suggestions}
    assert ("订单号", "order_id") in pairs
    assert ("订单编号", "order_id") not in pairs
    assert ("客户编号", "customer_id") in pairs
    assert ("客户名称", "customer_name") in pairs
    assert ("手机号", "mobile") in pairs
    assert ("订单金额", "total_amount") in pairs
    assert ("实付金额", "paid_amount") in pairs
    assert ("订单状态", "order_status") in pairs
    assert ("创建时间", "created_at") in pairs
    assert ("商品编码", "sku") in pairs


@pytest.mark.parametrize(
    ("source_column", "target_column"),
    [
        ("订单号", "order_id"),
        ("订单编号", "order_no"),
        ("客户编号", "customer_id"),
        ("客户名称", "customer_name"),
        ("手机号", "mobile"),
        ("订单金额", "total_amount"),
        ("实付金额", "paid_amount"),
        ("订单状态", "order_status"),
        ("创建时间", "created_at"),
        ("商品编码", "sku"),
    ],
)
def test_chinese_mapping_golden_pair(source_column: str, target_column: str):
    source = pd.DataFrame({source_column: ["A001"]})
    target = pd.DataFrame({target_column: ["A001"]})
    left, right = LoadedTable(source, "source.csv", None), LoadedTable(target, "target.csv", None)
    suggestions = suggest_column_mappings(profile_dataframe(left), profile_dataframe(right), source, target)
    assert (source_column, target_column) in {
        (item.source_column, item.target_column) for item in suggestions
    }


def test_string_comparison_is_case_sensitive_by_default():
    from reconcile_skill.normalizers import values_equal

    equal, kind = values_equal("ABC001", "abc001", "string", None, ["trim", "collapse_whitespace"])
    assert equal is False
    assert kind == "VALUE_MISMATCH"


def test_mapping_collision_is_blocked():
    suggestions = [
        MappingSuggestion(source_column="a", target_column="x", confidence=0.9, mapping_type="string"),
        MappingSuggestion(source_column="b", target_column="x", confidence=0.9, mapping_type="string"),
    ]
    try:
        rules_from_suggestions(suggestions)
    except ReconciliationError as exc:
        assert exc.code == "MAPPING_COLLISION"
    else:
        raise AssertionError("many-to-one mapping must not execute")


def test_gb18030_csv_fallback(tmp_path: Path):
    path = tmp_path / "cn.csv"
    path.write_bytes("订单号,金额\nA001,100\n".encode("gb18030"))
    table = load_table(path)
    assert list(table.dataframe.columns) == ["订单号", "金额"]


def test_formula_like_strings_are_safe_and_negative_numbers_survive(tmp_path: Path):
    assert sanitize_excel_cell("=1+1").startswith("'")
    assert sanitize_excel_cell("+SUM(A1:A2)").startswith("'")
    assert sanitize_excel_cell("@SUM(1,1)").startswith("'")
    assert sanitize_excel_cell("=1+1").startswith("'")
    assert sanitize_excel_cell("-123.45") == "-123.45"
    result = __import__("reconcile_skill.reconciler", fromlist=["reconcile"]).reconcile(
        pd.DataFrame(
            {
                "id": ["A001", "A002", "A003", "A004"],
                "value": [
                    "=1+1",
                    '=HYPERLINK("https://example.com")',
                    "@SUM(1,1)",
                    "+SUM(1,1)",
                ],
            }
        ),
        pd.DataFrame({"id": ["A001", "A002", "A003", "A004"], "value": ["safe"] * 4}),
        source_key="id",
        target_key="id",
        rules=[CompareRule(source_column="value", target_column="value", rule_type="string")],
    )
    report, _ = write_report(result, tmp_path, source_key="id", target_key="id")
    workbook = load_workbook(report, data_only=False, read_only=True)
    assert all(cell.data_type != "f" for sheet in workbook.worksheets for row in sheet.iter_rows() for cell in row)
