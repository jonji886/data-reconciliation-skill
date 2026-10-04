from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from reconcile_skill.errors import ReconciliationError
from reconcile_skill.config import load_mapping_yaml
from reconcile_skill.loaders import load_table
from reconcile_skill.models import CompareRule
from reconcile_skill.reconciler import reconcile
from reconcile_skill.report import write_report


def test_csv_and_multi_sheet_selection(tmp_path: Path):
    csv_path = tmp_path / "data.csv"
    csv_path.write_text("id;value\nA001;1\n", encoding="utf-8")
    table = load_table(csv_path)
    assert list(table.dataframe.columns) == ["id", "value"]

    xlsx_path = tmp_path / "multi.xlsx"
    with pd.ExcelWriter(xlsx_path) as writer:
        pd.DataFrame({"id": ["A"]}).to_excel(writer, sheet_name="one", index=False)
        pd.DataFrame({"id": ["B"]}).to_excel(writer, sheet_name="two", index=False)
    with pytest.raises(ReconciliationError) as error:
        load_table(xlsx_path)
    assert error.value.code == "SCHEMA_ERROR"
    assert load_table(xlsx_path, sheet_name="two").dataframe.iloc[0, 0] == "B"


def test_report_has_required_sheets(tmp_path: Path):
    result = reconcile(
        pd.DataFrame({"id": ["A"]}),
        pd.DataFrame({"id": ["A"]}),
        source_key="id",
        target_key="id",
        rules=[],
    )
    report, mapping = write_report(result, tmp_path, source_key="id", target_key="id")
    assert report.exists() and mapping.exists()
    loaded_mapping = load_mapping_yaml(mapping)
    assert loaded_mapping.key == {"source": "id", "target": "id", "case_sensitive": True}
    workbook = load_workbook(report, read_only=True)
    assert set(workbook.sheetnames) == {
        "Summary", "Mapping", "Missing_In_Source", "Missing_In_Target", "Duplicate_Source",
        "Duplicate_Target", "Value_Mismatch", "Diagnostics",
    }
    summary_labels = {row[0].value for row in workbook["Summary"].iter_rows() if row[0].value}
    assert "Matched Unique Keys" in summary_labels
    assert "Duplicate Source Keys" in summary_labels
    assert "Duplicate Source Rows" in summary_labels
